from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "outputs" / "user-guides"
SCREENSHOT_DIR = OUT_DIR / "screenshots"
LOGO = ROOT / "wp-content" / "themes" / "gsf-hub-sunrise" / "assets" / "images" / "gsf-logo-updated.png"

END_USER_OUTPUT = OUT_DIR / "GSF-Hub-End-User-Guide.docx"
ADMIN_OUTPUT = OUT_DIR / "GSF-Hub-Administrator-Guide.docx"
STAFF_OUTPUT = OUT_DIR / "GSF-Hub-Staff-Content-Operations-Guide.docx"

VERSION_DATE = date(2026, 7, 28).strftime("%B %-d, %Y")

# compact_reference_guide tokens, with named GSF brand overrides.
FONT = "Calibri"
INK = RGBColor(32, 45, 52)
MUTED = RGBColor(83, 105, 113)
NAVY = RGBColor(0, 82, 105)
TEAL = RGBColor(28, 166, 166)
GOLD = RGBColor(205, 145, 35)
WHITE = RGBColor(255, 255, 255)
LIGHT_TEAL = "EAF7F6"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F4F6F9"
MID_GRAY = "D9E2E7"
CAUTION = "FFF4D9"
RISK = "FDECEC"
TABLE_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120


def set_run_font(run, *, size=None, bold=None, italic=None, color=None, name=FONT):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def set_paragraph_tokens(paragraph, *, before=0, after=6, line=1.25, keep_next=False):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    fmt.keep_with_next = keep_next


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, *, color="C8D5DA", size=6):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), str(size))
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def set_table_geometry(table, widths_dxa, indent_dxa=TABLE_INDENT_DXA):
    total = sum(widths_dxa)
    table.autofit = False
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            width = widths_dxa[min(index, len(widths_dxa) - 1)]
            tc_w = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                cell._tc.get_or_add_tcPr().append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")


def mark_header_row(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    cant_split.set(qn("w:val"), "true")
    tr_pr.append(cant_split)


def configure_list_numbering(doc):
    numbering = doc.part.numbering_part.element
    for style_name, suffix in (("List Bullet", "space"), ("List Number", "tab")):
        style_num_pr = doc.styles[style_name]._element.pPr.numPr
        if style_num_pr is None or style_num_pr.numId is None:
            continue
        num_id = int(style_num_pr.numId.val)
        num = next(
            (
                node for node in numbering.findall(qn("w:num"))
                if int(node.get(qn("w:numId"))) == num_id
            ),
            None,
        )
        if num is None:
            continue
        abstract_num_id = int(num.find(qn("w:abstractNumId")).get(qn("w:val")))
        abstract = next(
            (
                node for node in numbering.findall(qn("w:abstractNum"))
                if int(node.get(qn("w:abstractNumId"))) == abstract_num_id
            ),
            None,
        )
        if abstract is None:
            continue
        level = next(
            (
                node for node in abstract.findall(qn("w:lvl"))
                if int(node.get(qn("w:ilvl"), "0")) == 0
            ),
            None,
        )
        if level is None:
            continue
        suff = level.find(qn("w:suff"))
        if suff is None:
            suff = OxmlElement("w:suff")
            level.append(suff)
        suff.set(qn("w:val"), suffix)

        p_pr = level.find(qn("w:pPr"))
        if p_pr is None:
            p_pr = OxmlElement("w:pPr")
            level.append(p_pr)
        ind = p_pr.find(qn("w:ind"))
        if ind is None:
            ind = OxmlElement("w:ind")
            p_pr.append(ind)
        ind.set(qn("w:left"), "540")
        ind.set(qn("w:hanging"), "270")

        tabs = p_pr.find(qn("w:tabs"))
        if tabs is None:
            tabs = OxmlElement("w:tabs")
            p_pr.insert(0, tabs)
        for child in list(tabs):
            tabs.remove(child)
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "num")
        tab.set(qn("w:pos"), "540")
        tabs.append(tab)


def set_picture_alt(run, description):
    drawings = run._element.xpath(".//wp:docPr")
    for drawing in drawings:
        drawing.set("descr", description)
        drawing.set("title", description)


def add_page_field(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (begin, instr, separate, text, end):
        run._r.append(node)
    set_run_font(run, size=9, color=MUTED)


def configure_styles(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    title = doc.styles["Title"]
    title.font.name = FONT
    title._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    title._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    title.font.size = Pt(30)
    title.font.bold = True
    title.font.color.rgb = NAVY
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(8)

    subtitle = doc.styles["Subtitle"]
    subtitle.font.name = FONT
    subtitle._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    subtitle._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    subtitle.font.size = Pt(14)
    subtitle.font.color.rgb = MUTED
    subtitle.paragraph_format.space_after = Pt(18)

    for name, size, color, before, after in (
        ("Heading 1", 16, NAVY, 18, 10),
        ("Heading 2", 13, NAVY, 14, 7),
        ("Heading 3", 12, RGBColor(31, 77, 120), 10, 5),
    ):
        style = doc.styles[name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        style = doc.styles[name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style.font.size = Pt(11)
        style.font.color.rgb = INK
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25

    configure_list_numbering(doc)

    caption = doc.styles["Caption"]
    caption.font.name = FONT
    caption._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    caption._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    caption.font.size = Pt(9.5)
    caption.font.italic = True
    caption.font.color.rgb = MUTED
    caption.paragraph_format.space_before = Pt(0)
    caption.paragraph_format.space_after = Pt(8)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER


def set_header_footer(doc, short_title):
    for section in doc.sections:
        header = section.header
        paragraph = header.paragraphs[0]
        paragraph.text = ""
        set_paragraph_tokens(paragraph, after=0, line=1.0)
        run = paragraph.add_run(f"GSF Hub  |  {short_title}")
        set_run_font(run, size=9, bold=True, color=MUTED)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT

        footer = section.footer
        footer_p = footer.paragraphs[0]
        footer_p.text = ""
        footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_paragraph_tokens(footer_p, after=0, line=1.0)
        run = footer_p.add_run("GSF Hub guide  |  ")
        set_run_font(run, size=9, color=MUTED)
        add_page_field(footer_p)


def add_para(doc, text="", *, bold_prefix=None, italic=False, color=INK, after=6, keep_next=False):
    paragraph = doc.add_paragraph()
    set_paragraph_tokens(paragraph, after=after, keep_next=keep_next)
    if bold_prefix and text.startswith(bold_prefix):
        prefix = paragraph.add_run(bold_prefix)
        set_run_font(prefix, size=11, bold=True, color=color)
        rest = paragraph.add_run(text[len(bold_prefix):])
        set_run_font(rest, size=11, italic=italic, color=color)
    else:
        run = paragraph.add_run(text)
        set_run_font(run, size=11, italic=italic, color=color)
    return paragraph


def add_heading(doc, text, level=1):
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    run = paragraph.add_run(text)
    set_run_font(
        run,
        size={1: 16, 2: 13, 3: 12}[level],
        bold=True,
        color=NAVY if level < 3 else RGBColor(31, 77, 120),
    )
    return paragraph


def add_bullet(doc, text, *, bold_prefix=None):
    paragraph = doc.add_paragraph(style="List Bullet")
    if bold_prefix and text.startswith(bold_prefix):
        run = paragraph.add_run(bold_prefix)
        set_run_font(run, size=11, bold=True, color=INK)
        run = paragraph.add_run(text[len(bold_prefix):])
        set_run_font(run, size=11, color=INK)
    else:
        run = paragraph.add_run(text)
        set_run_font(run, size=11, color=INK)
    return paragraph


def add_number(doc, text, *, bold_prefix=None):
    previous_num_id = None
    if doc.paragraphs:
        previous = doc.paragraphs[-1]
        if previous.style and previous.style.name == "List Number":
            num_pr = previous._p.pPr.numPr if previous._p.pPr is not None else None
            if num_pr is not None and num_pr.numId is not None:
                previous_num_id = int(num_pr.numId.val)

    paragraph = doc.add_paragraph(style="List Number")
    if previous_num_id is None:
        numbering = doc.part.numbering_part.element
        style_num_pr = doc.styles["List Number"]._element.pPr.numPr
        base_num_id = int(style_num_pr.numId.val)
        base_num = next(
            node for node in numbering.findall(qn("w:num"))
            if int(node.get(qn("w:numId"))) == base_num_id
        )
        abstract_num_id = int(base_num.find(qn("w:abstractNumId")).get(qn("w:val")))
        existing_ids = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
        previous_num_id = max(existing_ids, default=0) + 1
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(previous_num_id))
        abstract = OxmlElement("w:abstractNumId")
        abstract.set(qn("w:val"), str(abstract_num_id))
        num.append(abstract)
        override = OxmlElement("w:lvlOverride")
        override.set(qn("w:ilvl"), "0")
        start = OxmlElement("w:startOverride")
        start.set(qn("w:val"), "1")
        override.append(start)
        num.append(override)
        numbering.append(num)

    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.get_or_add_numPr()
    ilvl = num_pr.get_or_add_ilvl()
    ilvl.val = 0
    num_id = num_pr.get_or_add_numId()
    num_id.val = previous_num_id

    if bold_prefix and text.startswith(bold_prefix):
        run = paragraph.add_run(bold_prefix)
        set_run_font(run, size=11, bold=True, color=INK)
        run = paragraph.add_run(text[len(bold_prefix):])
        set_run_font(run, size=11, color=INK)
    else:
        run = paragraph.add_run(text)
        set_run_font(run, size=11, color=INK)
    return paragraph


def add_note(doc, title, text, *, kind="note"):
    fill = {"note": LIGHT_TEAL, "caution": CAUTION, "risk": RISK}.get(kind, LIGHT_GRAY)
    border = {"note": "8BCFCB", "caution": "E3BE62", "risk": "E3A5A5"}.get(kind, "C8D5DA")
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    prevent_row_split(table.rows[0])
    set_table_geometry(table, [TABLE_WIDTH_DXA], indent_dxa=160)
    set_table_borders(table, color=border, size=8)
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_margins(cell, top=120, bottom=120, start=160, end=160)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    set_paragraph_tokens(p, after=3, line=1.15)
    run = p.add_run(title)
    set_run_font(run, size=11, bold=True, color=NAVY)
    p2 = cell.add_paragraph()
    set_paragraph_tokens(p2, after=0, line=1.2)
    run = p2.add_run(text)
    set_run_font(run, size=10.5, color=INK)
    spacer = doc.add_paragraph()
    set_paragraph_tokens(spacer, after=2, line=1.0)
    return table


def add_table(doc, headers, rows, widths_dxa, *, header_fill=LIGHT_BLUE, font_size=10):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_geometry(table, widths_dxa)
    set_table_borders(table)
    mark_header_row(table.rows[0])
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        set_cell_shading(cell, header_fill)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        set_paragraph_tokens(p, after=0, line=1.15)
        run = p.add_run(header)
        set_run_font(run, size=font_size, bold=True, color=NAVY)
    for row_values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row_values):
            cell = cells[index]
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            set_paragraph_tokens(p, after=0, line=1.15)
            run = p.add_run(str(value))
            set_run_font(run, size=font_size, color=INK)
    spacer = doc.add_paragraph()
    set_paragraph_tokens(spacer, after=2, line=1.0)
    return table


def add_figure(doc, filename, caption, *, width=5.95):
    path = SCREENSHOT_DIR / filename
    if not path.exists():
        add_note(doc, "Screenshot unavailable", f"The expected test-site capture {filename} was not found.", kind="caution")
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_tokens(p, before=4, after=3, line=1.0, keep_next=True)
    run = p.add_run()
    run.add_picture(str(path), width=Inches(width))
    set_picture_alt(run, caption)
    cap = doc.add_paragraph(style="Caption")
    cap_run = cap.add_run(caption)
    set_run_font(cap_run, size=9.5, italic=True, color=MUTED)


def add_cover(doc, title, subtitle, audience, scope, short_title):
    set_header_footer(doc, short_title)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_tokens(p, before=16, after=22, line=1.0)
    if LOGO.exists():
        run = p.add_run()
        run.add_picture(str(LOGO), width=Inches(2.0))
        set_picture_alt(run, "GSF Hub logo")

    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_tokens(kicker, after=10, line=1.0)
    run = kicker.add_run("PRACTICAL USER GUIDE")
    set_run_font(run, size=10.5, bold=True, color=TEAL)

    title_p = doc.add_paragraph(style="Title")
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run(title)
    set_run_font(title_run, size=30, bold=True, color=NAVY)

    sub_p = doc.add_paragraph(style="Subtitle")
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = sub_p.add_run(subtitle)
    set_run_font(sub_run, size=14, color=MUTED)

    add_table(
        doc,
        ["Guide detail", "Description"],
        [
            ["Audience", audience],
            ["Scope", scope],
            ["Tested against", "GSF Hub local test site at localhost:8090"],
            ["Version date", VERSION_DATE],
        ],
        [2100, 7260],
    )
    add_note(
        doc,
        "How to use this guide",
        "Follow the numbered procedures in order. Screenshots show the current local test build; labels may move slightly after WordPress, theme, or plugin updates.",
    )
    doc.add_page_break()


def add_contents(doc, entries):
    add_heading(doc, "Contents", 1)
    for index, entry in enumerate(entries, start=1):
        p = doc.add_paragraph()
        set_paragraph_tokens(p, after=5, line=1.15)
        run = p.add_run(f"{index}.  {entry}")
        set_run_font(run, size=11, bold=index <= 2, color=NAVY if index <= 2 else INK)
    doc.add_page_break()


def begin_section(doc, title, intro=None):
    add_heading(doc, title, 1)
    if intro:
        add_para(doc, intro, color=MUTED, after=10)


def build_end_user_guide():
    doc = Document()
    configure_styles(doc)
    doc.core_properties.title = "GSF Hub End-User Guide"
    doc.core_properties.subject = "Public and signed-in member workflows"
    doc.core_properties.author = "GSF Hub"

    add_cover(
        doc,
        "GSF Hub End-User Guide",
        "Navigate the platform, create an account, learn, find resources, use data, and join the community.",
        "Visitors, registered members, learners, partners, and expert-roster applicants.",
        "Public navigation plus signed-in member workflows. WordPress administration is outside this guide.",
        "End-User Guide",
    )
    add_contents(
        doc,
        [
            "Quick start and navigation",
            "Create an account, sign in, and manage your profile",
            "Use Learn and earn a certificate",
            "Find resources and case studies",
            "Explore the Data Centre and Expert Directory",
            "Join forum discussions",
            "Subscribe to updates",
            "Troubleshooting and quick reference",
        ],
    )

    add_heading(doc, "1. Quick start and navigation", 1)
    add_para(
        doc,
        "The GSF Hub brings learning, practical tools, project stories, monitoring data, experts, and peer exchange into one platform. Most public content can be browsed without an account; learning enrollment, profile changes, and forum participation require sign-in.",
    )
    add_figure(doc, "end-01-home.png", "Figure 1. GSF Hub homepage and primary navigation on the local test site.")
    add_heading(doc, "Use the main menu", 2)
    add_table(
        doc,
        ["Menu", "What you can do"],
        [
            ["Learn", "Browse courses, enroll, launch course content, view progress, submit feedback, and access certificates."],
            ["Resources", "Search and filter templates, toolkits, reports, videos, and field-ready materials."],
            ["Case Studies", "Read outcome stories and filter them by type, focus area, and country or region."],
            ["Data Centre", "View project, participation, geographic, indicator, and Ecoequity reporting."],
            ["Directory", "Find practitioners and partner organizations by name, profile text, and country."],
            ["Forum", "Browse community spaces; sign in to create topics, reply, and subscribe."],
        ],
        [1650, 7710],
    )
    add_heading(doc, "Search from the homepage", 2)
    add_number(doc, "Open the homepage.")
    add_number(doc, "Enter a keyword in Search the site.")
    add_number(doc, "Select Search, then open the most relevant result.")
    add_note(doc, "Tip", "Use a short, specific phrase such as gender toolkit, Saint Lucia, or project management. Use each section's filters when you already know the content type.")

    begin_section(
        doc,
        "2. Create an account, sign in, and manage your profile",
        "Create one account and use it for learning, forum participation, expert-roster activity, and profile settings.",
    )
    add_heading(doc, "Create an account", 2)
    add_figure(doc, "end-10-register.png", "Figure 2. Member registration form.")
    for step in (
        "Select Sign in in the site header, then select Register.",
        "Enter a username, first name, last name, email address, and a strong password.",
        "Complete the optional profile information: gender, age, organization type, country, and stakeholder groups.",
        "Select Register. Follow any activation or approval instructions sent by email.",
        "Return to Sign in and use your username or email address and password.",
    ):
        add_number(doc, step)
    add_note(
        doc,
        "Privacy reminder",
        "Only provide profile and stakeholder information you are comfortable storing in your GSF Hub account. Review Privacy settings from your Account page after signing in.",
        kind="caution",
    )
    add_heading(doc, "Sign in or reset your password", 2)
    add_figure(doc, "end-09-sign-in.png", "Figure 3. Member sign-in page, including registration and password-reset links.")
    add_bullet(doc, "Use Forgot your password? if you cannot sign in.")
    add_bullet(doc, "Use Stay logged in only on a private device.")
    add_bullet(doc, "If a protected page sent you to Sign in, the site should return you to that page after successful login.")
    add_heading(doc, "Use your Account page", 2)
    add_figure(doc, "end-11-account.png", "Figure 4. Signed-in member Account page and Quick Access cards.")
    add_bullet(doc, "Update your first name, last name, or email address under Account.")
    add_bullet(doc, "Use Change Password to replace your password.")
    add_bullet(doc, "Use Privacy to control directory visibility where available.")
    add_bullet(doc, "Use the Quick Access cards to open Learn, the Expert Roster portal, or the Expert Directory.")
    add_note(doc, "Do not share accounts", "Course progress, feedback, certificates, subscriptions, and forum activity are tied to the signed-in account.")

    begin_section(
        doc,
        "3. Use Learn and earn a certificate",
        "The Learn area combines a course catalogue with a personal progress dashboard.",
    )
    add_figure(doc, "end-12-learning-signed-in.png", "Figure 5. Signed-in Learn catalogue with Enroll buttons and the learner dashboard.")
    add_heading(doc, "Enroll and start a course", 2)
    for step in (
        "Open Learn and choose a course title.",
        "Select Enroll. The page changes from Enroll to Status: Enrolled.",
        "Open or continue the embedded course content.",
        "Complete every required screen, activity, and assessment. Course progress is saved to your account.",
        "Return to Learn to review your course status.",
    ):
        add_number(doc, step)
    add_figure(doc, "end-16-course-enrolled.png", "Figure 6. Enrolled course page showing course status and the feedback step.")
    add_heading(doc, "Complete feedback and access the certificate", 2)
    add_bullet(doc, "The course package must report Completed or Passed before feedback becomes available.")
    add_bullet(doc, "After completion, select Submit feedback survey and submit the required responses.")
    add_bullet(doc, "The certificate is issued only after verified course completion and feedback.")
    add_bullet(doc, "Use Print or Save PDF on the certificate card to retain a copy.")
    add_note(
        doc,
        "If progress does not update",
        "Keep the course page open while studying, allow the final completion screen to finish saving, then refresh the course page. If the problem remains, record the course name, date, and last completed section before contacting support.",
        kind="caution",
    )

    begin_section(doc, "4. Find resources and case studies")
    add_heading(doc, "Search the Resource Library", 2)
    add_figure(doc, "end-04-resources.png", "Figure 7. Resource Library search and filters.")
    for step in (
        "Open Resources.",
        "Enter a keyword, or choose a resource type, topic, and country or region.",
        "Select Apply filters.",
        "Open a resource card to review its description and download or external link.",
        "Clear the fields or use Reset when you want to start a new search.",
    ):
        add_number(doc, step)
    add_heading(doc, "Browse case studies", 2)
    add_figure(doc, "end-05-case-studies.png", "Figure 8. Case Studies page with search and filter controls.")
    add_bullet(doc, "Search by title or keyword.")
    add_bullet(doc, "Narrow results by case study type, focus area, and country or region.")
    add_bullet(doc, "Open a result to review the context, actions, outcomes, partners, and reusable lessons.")

    begin_section(doc, "5. Explore the Data Centre and Expert Directory")
    add_heading(doc, "Read the Data Centre", 2)
    add_figure(doc, "end-06-data-centre.png", "Figure 9. Public Data Centre storytelling dashboard.")
    add_bullet(doc, "Use summary cards for the current number of projects, territories, participants, and course completions.")
    add_bullet(doc, "Scroll through participation, geographic, project-status, indicator, and Ecoequity sections.")
    add_bullet(doc, "Treat dashboard figures as the current published view; contact GSF staff if a figure needs clarification.")
    add_heading(doc, "Find an expert", 2)
    add_figure(doc, "end-07-directory.png", "Figure 10. Expert Directory search and country filter.")
    for step in (
        "Open Directory.",
        "Search by name, organization, role, or profile text.",
        "Choose a country when you need a regional match.",
        "Open a profile to review public information and areas of expertise.",
    ):
        add_number(doc, step)

    begin_section(doc, "6. Join forum discussions")
    add_figure(doc, "end-14-forum-signed-in.png", "Figure 11. Signed-in Forum page with community spaces and activity counts.")
    add_heading(doc, "Create a topic", 2)
    for step in (
        "Sign in and open Forum.",
        "Choose the forum space that best matches your question or update.",
        "Scroll to Create New Topic.",
        "Enter a clear title, write the topic, and add relevant tags.",
        "Optionally select Notify me of follow-up replies via email.",
        "Select Submit once. Your topic appears in that forum.",
    ):
        add_number(doc, step)
    add_figure(doc, "end-15-forum-new-topic.png", "Figure 12. New-topic form in a community forum.")
    add_heading(doc, "Reply and subscribe", 2)
    add_bullet(doc, "Open a topic and use the reply form to contribute.")
    add_bullet(doc, "Use Subscribe on a forum or topic when you want email updates.")
    add_bullet(doc, "Keep posts relevant, respectful, and free of confidential personal or project information.")
    add_note(doc, "Forum etiquette", "Use descriptive titles, acknowledge sources, and avoid posting the same question in multiple spaces.")

    begin_section(doc, "7. Subscribe to updates")
    add_para(doc, "The Mailing List form appears in the site footer.")
    for step in (
        "Scroll to Mailing List.",
        "Enter your email address.",
        "Select Subscribe.",
        "Complete any confirmation step sent by email.",
    ):
        add_number(doc, step)
    add_note(doc, "Subscription is separate from membership", "A GSF Hub account does not automatically subscribe you to newsletters, and a newsletter subscription does not create a member account.")

    begin_section(doc, "8. Troubleshooting and quick reference")
    add_table(
        doc,
        ["Problem", "What to try"],
        [
            ["Cannot sign in", "Check username or email, use Forgot your password?, and check junk mail for the reset message."],
            ["No course content", "Confirm you are signed in and enrolled, then refresh the course page."],
            ["Course still in progress", "Finish the course's final screen and allow it to save before refreshing."],
            ["Feedback is locked", "The course must first report Completed or Passed."],
            ["No certificate", "Complete the course, submit feedback, then refresh and open the certificate card."],
            ["No filter results", "Clear one or more filters, shorten the keyword, or reset the form."],
            ["Cannot post in Forum", "Confirm you are signed in and that your account has forum participation access."],
        ],
        [2300, 7060],
    )
    add_heading(doc, "Before contacting support", 2)
    add_bullet(doc, "Record the page or course title and the time the problem occurred.")
    add_bullet(doc, "Take a screenshot of the visible message without exposing your password.")
    add_bullet(doc, "Note the browser and device you were using.")

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc.save(END_USER_OUTPUT)
    return END_USER_OUTPUT


def build_admin_guide():
    doc = Document()
    configure_styles(doc)
    doc.core_properties.title = "GSF Hub Administrator Guide"
    doc.core_properties.subject = "WordPress site administration, governance, and onboarding"
    doc.core_properties.author = "GSF Hub"

    add_cover(
        doc,
        "GSF Hub Administrator Guide",
        "Manage people, permissions, platform configuration, updates, backups, and operational health.",
        "WordPress administrators, platform owners, and technical administrators.",
        "Site-wide administration. Day-to-day content production is summarized here and covered in the Staff Content Operations Guide.",
        "Administrator Guide",
    )
    add_contents(
        doc,
        [
            "Administrator responsibilities and safe operating pattern",
            "Dashboard and routine checks",
            "Users, registration, roles, and expert applications",
            "Plugins, themes, updates, and site configuration",
            "Backups, migration, email, and site health",
            "First-time end-user display assessment",
            "Incident response and handover checklist",
        ],
    )

    add_heading(doc, "1. Administrator responsibilities and safe operating pattern", 1)
    add_para(
        doc,
        "Administrators have full control of WordPress, plugins, themes, users, roles, settings, and all GSF content areas. Use the administrator role only when site-wide authority is required.",
    )
    add_note(
        doc,
        "Administrator standard",
        "Back up before structural changes, make one change set at a time, test on localhost or staging, verify the public result, and document anything that changes access or data behavior.",
        kind="caution",
    )
    add_heading(doc, "Sign in", 2)
    add_number(doc, "Open /wp-admin/ on the correct environment.")
    add_number(doc, "Use your assigned administrator account; do not share credentials.")
    add_number(doc, "Confirm the site name and environment before changing anything.")
    add_number(doc, "Use the WordPress toolbar to view the public site after changes.")
    add_heading(doc, "Before any high-impact change", 2)
    for item in (
        "Confirm a recent backup exists and can be restored.",
        "Record the current plugin or theme version.",
        "Check whether the change affects registration, login, Learn, Data Centre, multilingual content, or forums.",
        "Plan a rollback before updating, importing, changing roles, or editing settings.",
    ):
        add_bullet(doc, item)

    begin_section(doc, "2. Dashboard and routine checks")
    add_figure(doc, "admin-01-dashboard.png", "Figure 1. WordPress administrator dashboard and GSF Hub menu groups.")
    add_heading(doc, "Daily or publishing-day checks", 2)
    add_bullet(doc, "Review notices for failed background actions, mail failures, or site-health warnings.")
    add_bullet(doc, "Check new users, pending expert applications, and comments or forum activity.")
    add_bullet(doc, "Verify newly published courses, resources, case studies, and data on the public site.")
    add_heading(doc, "Weekly checks", 2)
    add_bullet(doc, "Review available updates without applying them automatically to production.")
    add_bullet(doc, "Review Site Health and email delivery.")
    add_bullet(doc, "Confirm backups are recent and stored outside the web root.")
    add_bullet(doc, "Review administrator accounts and remove access that is no longer required.")
    add_heading(doc, "Monthly checks", 2)
    add_bullet(doc, "Test registration, password reset, member login, one Learn enrollment, one forum post, and newsletter signup.")
    add_bullet(doc, "Review role assignments and Data Centre access.")
    add_bullet(doc, "Archive superseded exports and document platform changes.")

    begin_section(doc, "3. Users, registration, roles, and expert applications")
    add_heading(doc, "Create or update a user", 2)
    for step in (
        "Open Users > All Users or Add User.",
        "Search for an existing account before creating a duplicate.",
        "Set the minimum role needed for the person's work.",
        "Add a bbPress forum role only when forum participation or moderation is required.",
        "Save the account and ask the user to set or reset their own password.",
    ):
        add_number(doc, step)
    add_heading(doc, "GSF role guide", 2)
    add_table(
        doc,
        ["Role", "Recommended use"],
        [
            ["Subscriber / GSF Student", "Member or learner with front-end access only."],
            ["GSF Content Staff", "Resources, case studies, posts/pages, media, and Learn content where granted."],
            ["GSF MEAL Officer", "Data submission, verification, indicators, evidence, projects, and reports."],
            ["GSF CBF Staff / GSF Admin", "Broad Data Centre management; settings access depends on capability set."],
            ["bbPress Participant", "Create and edit own forum topics and replies."],
            ["bbPress Moderator", "Moderate topics and replies, including other users' content."],
            ["bbPress Keymaster", "Full forum administration."],
            ["Administrator", "Full platform control; reserve for site owners and technical administrators."],
        ],
        [2700, 6660],
    )
    add_note(
        doc,
        "Least privilege",
        "Content, Data Centre, Learn, newsletter, and forum duties use different capabilities. Assign only the roles needed; one staff member may require a WordPress role plus a forum role.",
    )
    add_heading(doc, "Review and change roles", 2)
    add_figure(doc, "admin-03-role-editor.png", "Figure 2. User Role Editor for reviewing role capabilities.")
    add_bullet(doc, "Prefer assigning an existing tested role over editing capabilities directly.")
    add_bullet(doc, "Never remove capabilities from Administrator without a tested recovery path.")
    add_bullet(doc, "After a role change, test with a non-administrator account in staging.")
    add_heading(doc, "Ultimate Member and expert applications", 2)
    add_bullet(doc, "Ultimate Member controls the front-end Login, Register, Account, member directory, email templates, and account status.")
    add_bullet(doc, "Users > Expert Applications supports the expert-roster review workflow.")
    add_bullet(doc, "Keep expert application status, directory profile status, and WordPress roles aligned.")
    add_bullet(doc, "The current configuration sends a welcome email for an activated account.")

    begin_section(doc, "4. Plugins, themes, updates, and site configuration")
    add_figure(doc, "admin-04-plugins.png", "Figure 3. Installed Plugins screen.")
    add_heading(doc, "Safe update sequence", 2)
    for step in (
        "Create or confirm a restorable backup.",
        "Review the changelog and compatibility for WordPress, PHP, theme, and dependent plugins.",
        "Apply the update on localhost or staging.",
        "Test homepage, Login, Register, Account, Learn, Resources, Case Studies, Data Centre, Directory, Forum, and email.",
        "Apply to production during an agreed maintenance window.",
        "Repeat the smoke test and record the result.",
    ):
        add_number(doc, step)
    add_note(
        doc,
        "Do not bulk-update blindly",
        "The GSF Hub combines Ultimate Member, Polylang, bbPress, Mailchimp, SCORM/Learn, Data Centre, Pods, and custom theme code. Update in small batches so a regression can be isolated.",
        kind="risk",
    )
    add_heading(doc, "Configuration ownership", 2)
    add_table(
        doc,
        ["Area", "Administrator menu"],
        [
            ["Site name, timezone, and base settings", "Settings > General"],
            ["Homepage and page publishing", "Pages and Appearance"],
            ["Navigation", "Appearance > Menus"],
            ["Language and translations", "Languages"],
            ["Registration and member accounts", "Ultimate Member"],
            ["Newsletter signup form and audience", "Mailchimp"],
            ["Outbound email delivery", "WP Mail SMTP"],
            ["Forums", "Forum and Settings > Forums"],
            ["Roles and capabilities", "Users > User Role Editor or Members"],
            ["Custom content models", "Pods Admin; change only with technical review"],
        ],
        [3600, 5760],
    )

    begin_section(doc, "5. Backups, migration, email, and site health")
    add_heading(doc, "Backups and migration", 2)
    add_bullet(doc, "Use All-in-One WP Migration or the approved transfer workflow for a full-site package.")
    add_bullet(doc, "Store backups outside the live WordPress directory and protect them as sensitive data.")
    add_bullet(doc, "Test restore procedures; a backup is not reliable until a restore has been demonstrated.")
    add_bullet(doc, "After migration, replace localhost URLs and test file downloads, SCORM packages, email, and multilingual links.")
    add_heading(doc, "Email", 2)
    add_bullet(doc, "WP Mail SMTP manages WordPress mail delivery and diagnostics.")
    add_bullet(doc, "The local test environment routes mail to Mailpit at localhost:8025.")
    add_bullet(doc, "Test welcome, password-reset, reminder, subscription, and notification messages after mail changes.")
    add_heading(doc, "Site Health", 2)
    add_bullet(doc, "Open Tools > Site Health and review critical issues first.")
    add_bullet(doc, "Treat public debug output, old plugins, failed cron jobs, and insecure transport as release blockers.")
    add_bullet(doc, "Do not paste secrets, full database exports, or member data into tickets or public chats.")

    begin_section(
        doc,
        "6. First-time end-user display assessment",
        "This section answers whether a first-user or first-login display can be added for front-end members.",
    )
    add_note(
        doc,
        "Finding",
        "Yes. The current site can support a one-time welcome display for end users. It is not currently implemented as a first-login-only experience.",
    )
    add_heading(doc, "What exists today", 2)
    add_bullet(doc, "Ultimate Member welcome email is enabled.")
    add_bullet(doc, "After sign-in, members land on /account/ and see Quick Access cards.")
    add_bullet(doc, "Those cards appear on normal account visits; they are not tracked as a one-time onboarding display.")
    add_bullet(doc, "The WordPress Welcome panel shown on the administrator dashboard is admin-only and does not serve front-end members.")
    add_heading(doc, "Recommended implementation", 2)
    add_para(
        doc,
        "Add a small GSF onboarding plugin or theme-independent module that redirects a newly activated member to /account/?welcome=1 or displays an accessible panel on the Account page. Store a user-meta flag such as gsf_onboarding_completed so it appears once per user.",
    )
    add_table(
        doc,
        ["Element", "Recommendation"],
        [
            ["Audience", "Subscriber, GSF Student, expert applicants, and other front-end roles; exclude administrators and staff by default."],
            ["Trigger", "First successful front-end login after activation, not every login."],
            ["Content", "Complete profile, open Learn, browse Resources, visit Forum, and review privacy."],
            ["Controls", "Start tour, Skip for now, and Do not show again; all keyboard accessible."],
            ["State", "Persist a per-user completion or dismissal flag; provide an admin-only reset for testing."],
            ["Analytics", "Track view, step completion, skip, and destination clicks without storing sensitive field values."],
            ["Fallback", "If scripting is blocked, render a normal Account-page welcome card with the same actions."],
        ],
        [2200, 7160],
    )
    add_heading(doc, "Acceptance checklist", 2)
    for item in (
        "Appears once for a brand-new approved front-end user.",
        "Does not interrupt password reset, email activation, course deep links, or expert application redirects.",
        "Does not appear for administrators or content staff unless explicitly enabled.",
        "Works on mobile, with keyboard-only navigation, and with a screen reader.",
        "Skip and completion state survive logout and a new browser session.",
        "An administrator can reset the flag on a test account.",
        "The user can still reach every destination later from Account > Quick Access.",
    ):
        add_bullet(doc, item)
    add_note(
        doc,
        "Implementation boundary",
        "This is a small custom development task rather than a WordPress core setting. Build and test it against Ultimate Member login/registration redirects and the existing Account page.",
        kind="caution",
    )

    begin_section(doc, "7. Incident response and handover checklist")
    add_heading(doc, "If a change causes a problem", 2)
    for step in (
        "Stop making further unrelated changes.",
        "Record the time, affected pages, user role, error message, and last successful state.",
        "Disable or roll back only the suspected change on staging first when possible.",
        "Restore from backup if the site or data cannot be safely recovered in place.",
        "Retest the affected workflow and the main smoke-test paths.",
        "Document the root cause and prevention action.",
    ):
        add_number(doc, step)
    add_heading(doc, "Administrator handover", 2)
    add_bullet(doc, "Named owners for hosting, domain, DNS, backups, WordPress, mail, Mailchimp, and Zoom.")
    add_bullet(doc, "Current administrator list and emergency access procedure.")
    add_bullet(doc, "Backup location, restore instructions, and latest successful restore test.")
    add_bullet(doc, "Plugin/theme inventory, custom-code locations, and update history.")
    add_bullet(doc, "Known issues, scheduled tasks, and support contacts.")

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc.save(ADMIN_OUTPUT)
    return ADMIN_OUTPUT


def build_staff_guide():
    doc = Document()
    configure_styles(doc)
    doc.core_properties.title = "GSF Hub Staff Content Operations Guide"
    doc.core_properties.subject = "Newsletter, Data Centre, Learn, Case Studies, Resources, and Forums"
    doc.core_properties.author = "GSF Hub"

    add_cover(
        doc,
        "GSF Hub Staff Content Operations Guide",
        "Operate Newsletter, Data Centre, Learn, Case Studies, Resources, and Forums.",
        "GSF content staff, MEAL officers, CBF staff, course managers, and forum moderators.",
        "Day-to-day content and data operations. Platform-level updates, plugins, themes, users, and settings belong to administrators.",
        "Staff Operations Guide",
    )
    add_contents(
        doc,
        [
            "Access and publishing standard",
            "Newsletter and mailing-list form",
            "Data Centre",
            "Learn",
            "Resources",
            "Case Studies",
            "Forums",
            "Quality assurance and escalation",
        ],
    )

    add_heading(doc, "1. Access and publishing standard", 1)
    add_para(
        doc,
        "Open /wp-admin/ and use the menu for your assigned role. The visible menu can differ by role; contact an administrator if a required area is missing.",
    )
    add_table(
        doc,
        ["Work area", "Typical access"],
        [
            ["Newsletter form", "Administrator or a role specifically granted Mailchimp settings access."],
            ["Data Centre", "GSF MEAL Officer, GSF CBF Staff, GSF Admin, or Administrator."],
            ["Learn", "GSF Content Staff with Learn capability, course manager, or Administrator."],
            ["Resources and Case Studies", "GSF Content Staff, Editor-equivalent role, or Administrator."],
            ["Forum participation", "bbPress Participant."],
            ["Forum moderation", "bbPress Moderator, Keymaster, or Administrator."],
        ],
        [2800, 6560],
    )
    add_note(
        doc,
        "If access is missing",
        "Do not borrow another person's account. Ask an administrator to assign the smallest role or capability set needed for your duties.",
        kind="caution",
    )
    add_heading(doc, "Standard publish cycle", 2)
    for step in (
        "Prepare the approved text, files, image, language, metadata, and owner.",
        "Create or edit the item in WordPress.",
        "Save as Draft and review the preview.",
        "Check links, downloads, spelling, accessibility, and filters.",
        "Publish only after content approval.",
        "Open the public page in a signed-out view and verify the result.",
    ):
        add_number(doc, step)
    add_heading(doc, "Multilingual content", 2)
    add_bullet(doc, "Select the correct language before publishing.")
    add_bullet(doc, "Connect translations using the Languages panel; do not overwrite one language with another.")
    add_bullet(doc, "Verify each language's title, body, metadata, file, and public URL.")

    begin_section(
        doc,
        "2. Newsletter and mailing-list form",
        "The current WordPress integration manages the website signup form and selected Mailchimp audience. Campaign creation and sending remain in Mailchimp.",
    )
    add_figure(doc, "staff-01-newsletter.png", "Figure 1. Mailchimp Form Settings in WordPress.")
    add_heading(doc, "Configure the website signup form", 2)
    for step in (
        "Open Mailchimp > Form Settings.",
        "Confirm the selected audience or list is GSF Hub.",
        "Review Form copy, button text, and any success or error messaging.",
        "Select only the audience fields the website form should collect.",
        "Review the Form preview.",
        "Save changes, then submit a test address from the public site footer.",
    ):
        add_number(doc, step)
    add_heading(doc, "Create and send a newsletter", 2)
    add_bullet(doc, "Create the campaign in the connected Mailchimp account, not in the current WordPress plugin.")
    add_bullet(doc, "Use the correct audience, subject, preview text, sender identity, and approved content.")
    add_bullet(doc, "Send a test to internal reviewers and check desktop, mobile, links, alt text, and unsubscribe.")
    add_bullet(doc, "Schedule or send only after approval; review delivery, opens, clicks, and unsubscribes in Mailchimp.")
    add_note(
        doc,
        "Local testing",
        "WordPress transactional email is routed to Mailpit at localhost:8025. Mailchimp campaign sending may still require the external Mailchimp account and should not be tested with real audiences from localhost.",
        kind="caution",
    )

    begin_section(
        doc,
        "3. Data Centre",
        "The Data Centre combines the PMF dashboard with indicators, results, disaggregation, organizations, projects, evidence, activities, Ecoequity, CSV tools, and settings.",
    )
    add_figure(doc, "staff-02-data-centre.png", "Figure 2. Data Centre dashboard and its operating menu.")
    add_heading(doc, "Recommended data-entry order", 2)
    for step in (
        "Confirm the organization exists.",
        "Create or update the project and its classification fields.",
        "Confirm the relevant PMF indicator exists.",
        "Add or update the indicator result for the correct reporting period.",
        "Add disaggregated rows that reconcile to the result total.",
        "Attach supporting evidence and activity records.",
        "Review the public dashboard and exported report.",
    ):
        add_number(doc, step)
    add_heading(doc, "Dashboard and indicators", 2)
    add_figure(doc, "staff-10-data-indicators.png", "Figure 3. PMF Indicators screen.")
    add_bullet(doc, "Dashboard summarizes indicator, result, project, evidence, and Ecoequity counts.")
    add_bullet(doc, "Indicators define the PMF framework; change definitions only with MEAL approval.")
    add_bullet(doc, "Indicator Results are the reported values. Check indicator, project, period, status, unit, and evidence.")
    add_heading(doc, "Disaggregation", 2)
    add_figure(doc, "staff-11-data-disaggregation.png", "Figure 4. Disaggregation editor and inclusion dashboard.")
    add_bullet(doc, "Use Result ID to connect each row to the correct indicator result.")
    add_bullet(doc, "Choose the category, enter a clear value label, and enter the numeric value.")
    add_bullet(doc, "Save edits and confirm that disaggregated values do not exceed or contradict the reported total.")
    add_note(doc, "Data integrity", "Do not guess missing values. Use the approved zero, not applicable, or no-data treatment for the indicator and reporting period.", kind="caution")
    add_heading(doc, "CSV import and export", 2)
    add_figure(doc, "staff-12-data-import-export.png", "Figure 5. CSV import templates and export reports.")
    for step in (
        "Download the matching import template.",
        "Keep the header names and required identifiers unchanged.",
        "Validate dates, numeric fields, codes, and duplicates before upload.",
        "Choose the correct Import type, select the CSV file, and import.",
        "Review the result message and spot-check the affected records.",
        "Export the related report and reconcile totals.",
    ):
        add_number(doc, step)
    add_note(
        doc,
        "Before a large import",
        "Ask an administrator to confirm a recent backup and run the import on localhost or staging first.",
        kind="risk",
    )

    begin_section(
        doc,
        "4. Learn",
        "Learn staff manage SCORM packages, course records, learner reminders, progress, and certificate appearance.",
    )
    add_figure(doc, "staff-03-learn.png", "Figure 6. Learner Reminders settings.")
    add_heading(doc, "Upload and publish a course", 2)
    add_figure(doc, "staff-09-upload-course.png", "Figure 7. Upload SCORM Package screen.")
    for step in (
        "Open Learn > Upload SCORM Package.",
        "Upload the approved SCORM ZIP package and wait for successful extraction.",
        "Open Learn > Add Course or the course record created by the upload.",
        "Enter the title, public description, duration, featured image, language, and SCORM package connection.",
        "Keep Issue certificate after completion and feedback enabled when a certificate is required.",
        "Save as Draft, preview, and test enrollment, launch, completion, feedback, and certificate.",
        "Publish after the complete learner path passes.",
    ):
        add_number(doc, step)
    add_heading(doc, "Monitor learner progress", 2)
    add_bullet(doc, "Filter by course and status: Enrolled, In progress, Completed, or Failed.")
    add_bullet(doc, "Investigate stalled records using the learner, course, last activity, and SCORM completion status.")
    add_bullet(doc, "Do not manually mark completion unless an approved correction procedure exists.")
    add_heading(doc, "Learner reminders", 2)
    add_bullet(doc, "Enable reminders only after outbound email is working.")
    add_bullet(doc, "Set an inactivity period and a clear subject and message.")
    add_bullet(doc, "Send a test reminder before enabling the schedule.")
    add_bullet(doc, "Check Schedule status after saving.")
    add_heading(doc, "Certificate settings", 2)
    add_figure(doc, "staff-14-learn-certificates.png", "Figure 8. Certificate text, logos, and brand colors.")
    add_bullet(doc, "Review title, introduction, completion line, issuer, and footer.")
    add_bullet(doc, "Use approved GSF, Canada, and Caribbean Biodiversity Fund logo URLs.")
    add_bullet(doc, "Confirm primary and accent colors and test Print or Save PDF.")
    add_note(doc, "Certificate gate", "A learner's certificate is issued only after verified course completion and submitted feedback.")

    begin_section(doc, "5. Resources")
    add_figure(doc, "staff-04-resources.png", "Figure 9. Resource Library list in WordPress.")
    add_heading(doc, "Add a resource", 2)
    add_figure(doc, "staff-07-add-resource.png", "Figure 10. Add New Resource editor and metadata fields.")
    for step in (
        "Open Resources > Add Resource.",
        "Enter the title and a concise description in the editor.",
        "Set Resource Type, Topic, Country/Region, Publication Year, and Source Organization.",
        "Attach the resource file and/or enter the approved external URL.",
        "Set a featured image and the correct language.",
        "Use Featured Resource only for approved homepage or promoted placement.",
        "Save as Draft and preview.",
        "Publish, then verify search, filters, file download, and external links on /resources/.",
    ):
        add_number(doc, step)
    add_note(doc, "Files", "Use descriptive filenames, an accessible document format, a reasonable file size, and the approved version. Replace superseded files rather than leaving users with ambiguous duplicates.")

    begin_section(doc, "6. Case Studies")
    add_figure(doc, "staff-05-case-studies.png", "Figure 11. Case Studies list in WordPress.")
    add_heading(doc, "Add a case study", 2)
    add_figure(doc, "staff-08-add-case-study.png", "Figure 12. Add Case Study editor.")
    for step in (
        "Open Case Studies > Add Case Study.",
        "Enter a specific outcome-focused title and the approved narrative.",
        "Set Case Study Type, Focus Area, Country/Region, and Implementation Date.",
        "Enter Lead Organization and Partners.",
        "Set a featured image and meaningful image alternative text in the Media Library.",
        "Use Featured Case Study only with editorial approval.",
        "Select the correct language and connect translations.",
        "Preview, publish, and verify the listing card, filters, and detail page.",
    ):
        add_number(doc, step)
    add_heading(doc, "Editorial checklist", 2)
    add_bullet(doc, "The title names the place, intervention, or outcome.")
    add_bullet(doc, "Claims and numbers have an approved source.")
    add_bullet(doc, "People are represented with consent and appropriate safeguarding.")
    add_bullet(doc, "The story identifies reusable lessons, not only activities.")
    add_bullet(doc, "Partner names, dates, and geography are consistent with Data Centre records.")

    begin_section(
        doc,
        "7. Forums",
        "Forum staff organize spaces and moderate topics and replies. Community members create discussions from the public site.",
    )
    add_figure(doc, "staff-06-forum.png", "Figure 13. Forums list and bbPress administration.")
    add_heading(doc, "Manage forum spaces", 2)
    add_bullet(doc, "Use Forums to edit the title, description, order, status, and visibility of each forum.")
    add_bullet(doc, "Avoid renaming or deleting a live forum without checking links and moving its topics.")
    add_bullet(doc, "Use closed or private status only when the access decision is approved.")
    add_heading(doc, "Moderate topics", 2)
    add_bullet(doc, "Filter by forum and search by title or author.")
    add_bullet(doc, "Open a topic to review content, author, tags, status, and forum assignment.")
    add_bullet(doc, "Move misplaced topics, edit tags, close resolved threads, or mark spam according to policy.")
    add_heading(doc, "Moderate replies", 2)
    add_bullet(doc, "Review the parent forum and topic before changing a reply.")
    add_bullet(doc, "Use Trash or Spam only when the moderation policy supports it.")
    add_bullet(doc, "Escalate threats, harassment, sensitive personal information, or safeguarding concerns immediately.")
    add_note(doc, "Forum authority", "GSF Content Staff plus bbPress Participant can participate but cannot fully moderate. Moderation requires bbPress Moderator, Keymaster, or Administrator.")

    begin_section(doc, "8. Quality assurance and escalation")
    add_heading(doc, "Public verification checklist", 2)
    for item in (
        "Correct title, language, date, owner, and featured image.",
        "No draft notes, localhost-only text, broken shortcodes, or placeholder content.",
        "Links and downloads open the intended destination.",
        "Search and filters return the new or updated item.",
        "Mobile layout is readable and controls remain usable.",
        "Images have meaningful alternative text; headings follow a logical order.",
        "Restricted data, private contact details, and internal files are not public.",
    ):
        add_bullet(doc, item)
    add_heading(doc, "Escalate to an administrator when", 2)
    add_bullet(doc, "A menu is missing or a permission blocks approved work.")
    add_bullet(doc, "A plugin, theme, update, import, backup, restore, user-role, or site-setting change is required.")
    add_bullet(doc, "Registration, login, email, SCORM launch, certificate generation, or public data is malfunctioning.")
    add_bullet(doc, "A bulk change or large CSV import could affect many records.")
    add_bullet(doc, "You suspect a security, privacy, or safeguarding incident.")

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc.save(STAFF_OUTPUT)
    return STAFF_OUTPUT


def main():
    outputs = [
        build_end_user_guide(),
        build_admin_guide(),
        build_staff_guide(),
    ]
    for output in outputs:
        print(output)


if __name__ == "__main__":
    main()
