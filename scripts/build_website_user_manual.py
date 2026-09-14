from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "outputs" / "GSF-Hub-Website-User-Manual.docx"
SCREENSHOT_DIR = ROOT / "outputs" / "manual-screenshots"

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
INK = RGBColor(32, 32, 32)
MUTED = RGBColor(90, 90, 90)
FILL_BLUE_GRAY = "E8EEF5"
FILL_LIGHT = "F4F6F9"
BORDER = "D8DEE8"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=BORDER):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        node = borders.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "6")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def set_run_font(run, size=None, bold=None, color=None, italic=None):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color
    if italic is not None:
        run.italic = italic


def set_para_format(paragraph, before=0, after=6, line=1.25):
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line


def add_paragraph(doc, text="", style=None, bold_prefix=None):
    paragraph = doc.add_paragraph(style=style)
    set_para_format(paragraph)
    if bold_prefix and text.startswith(bold_prefix):
        prefix = paragraph.add_run(bold_prefix)
        set_run_font(prefix, size=11, bold=True, color=INK)
        rest = paragraph.add_run(text[len(bold_prefix):])
        set_run_font(rest, size=11, color=INK)
    else:
        run = paragraph.add_run(text)
        set_run_font(run, size=11, color=INK)
    return paragraph


def add_heading(doc, text, level=1):
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    if level == 1:
        paragraph.paragraph_format.space_before = Pt(18)
        paragraph.paragraph_format.space_after = Pt(10)
    elif level == 2:
        paragraph.paragraph_format.space_before = Pt(14)
        paragraph.paragraph_format.space_after = Pt(7)
    else:
        paragraph.paragraph_format.space_before = Pt(10)
        paragraph.paragraph_format.space_after = Pt(5)
    run = paragraph.add_run(text)
    set_run_font(run, size={1: 16, 2: 13, 3: 12}.get(level, 11), bold=True, color=BLUE if level < 3 else DARK_BLUE)
    return paragraph


def add_bullet(doc, text):
    paragraph = doc.add_paragraph(style="List Bullet")
    set_para_format(paragraph, after=4, line=1.25)
    run = paragraph.add_run(text)
    set_run_font(run, size=11, color=INK)
    return paragraph


def add_number(doc, text):
    paragraph = doc.add_paragraph(style="List Number")
    set_para_format(paragraph, after=4, line=1.25)
    run = paragraph.add_run(text)
    set_run_font(run, size=11, color=INK)
    return paragraph


def add_note(doc, title, text):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(6.5)
    set_table_borders(table, color="C9D6E6")
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_shading(cell, FILL_LIGHT)
    set_cell_margins(cell, top=120, bottom=120, start=160, end=160)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    set_run_font(r, size=11, bold=True, color=DARK_BLUE)
    p2 = cell.add_paragraph()
    set_para_format(p2, after=0)
    r2 = p2.add_run(text)
    set_run_font(r2, size=10.5, color=INK)


def add_screenshot(doc, filename, caption):
    image_path = SCREENSHOT_DIR / filename
    if not image_path.exists():
        return

    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    run.add_picture(str(image_path), width=Inches(6.2))

    caption_p = doc.add_paragraph()
    caption_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_p.paragraph_format.space_before = Pt(0)
    caption_p.paragraph_format.space_after = Pt(8)
    caption_run = caption_p.add_run(caption)
    set_run_font(caption_run, size=9.5, italic=True, color=MUTED)


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    for index, width in enumerate(widths):
        table.columns[index].width = Inches(width)
    header_cells = table.rows[0].cells
    for index, header in enumerate(headers):
        header_cells[index].width = Inches(widths[index])
        set_cell_shading(header_cells[index], FILL_BLUE_GRAY)
        set_cell_margins(header_cells[index])
        header_cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = header_cells[index].paragraphs[0]
        set_para_format(p, after=0)
        r = p.add_run(header)
        set_run_font(r, size=10.5, bold=True, color=DARK_BLUE)
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].width = Inches(widths[index])
            set_cell_margins(cells[index])
            cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cells[index].paragraphs[0]
            set_para_format(p, after=0)
            r = p.add_run(value)
            set_run_font(r, size=10.5, color=INK)
    doc.add_paragraph()
    return table


def setup_styles(doc):
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)

    for style_name in ("List Bullet", "List Number"):
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(11)


def add_cover(doc):
    section = doc.sections[0]
    header = section.header.paragraphs[0]
    header.text = ""
    run = header.add_run("GSF Hub | Website User Manual")
    set_run_font(run, size=9, color=MUTED)
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    footer = section.footer.paragraphs[0]
    footer.text = ""
    run = footer.add_run("Prepared for GSF Hub users")
    set_run_font(run, size=9, color=MUTED)
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(44)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("GSF Hub")
    set_run_font(r, size=28, bold=True, color=DARK_BLUE)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(18)
    r = p.add_run("Website User Manual")
    set_run_font(r, size=20, bold=True, color=BLUE)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(20)
    r = p.add_run("A practical guide for navigating the homepage, Learn, Resources, Case Studies, Data Centre, Directory, Forum, search, filters, feedback, and certificates.")
    set_run_font(r, size=12, color=INK)

    add_table(
        doc,
        ["Field", "Details"],
        [
            ["Audience", "General website users, learners, partners, and staff who need to find content or complete courses."],
            ["Scope", "Public navigation and signed-in learner workflows. Administration is noted only where it affects user-facing content."],
            ["Version date", date(2026, 7, 1).strftime("%B %-d, %Y")],
        ],
        [1.55, 4.95],
    )

    add_note(
        doc,
        "How to use this manual",
        "Start with the quick navigation table, then jump to the section for the area you are using. Most pages follow the same pattern: search, filter, open a card or title, then use Reset to clear filters.",
    )
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_before = Pt(14)
    spacer.paragraph_format.space_after = Pt(2)


def build_manual():
    doc = Document()
    setup_styles(doc)
    add_cover(doc)

    add_heading(doc, "1. Quick Navigation", 1)
    add_paragraph(
        doc,
        "The main navigation appears at the top of the website. On smaller screens, open the menu button to show the same areas.",
    )
    add_table(
        doc,
        ["Area", "Use it for", "Common actions"],
        [
            ["Homepage", "Starting point for hub-wide search, featured content, key metrics, resources, courses, and events.", "Use Search, open quick links, review featured cards, or return home by selecting the GSF Hub logo."],
            ["Learn", "Access self-paced courses, SCORM course content, progress tracking, feedback, and certificates.", "Open a course, complete the course package, submit feedback, then view the certificate."],
            ["Resources", "Find toolkits, templates, publications, learning notes, multimedia resources, and other reference materials.", "Search by keyword, filter by Type, Topic, or Country / Region, then select View details."],
            ["Case Studies", "Browse examples, project stories, and implementation evidence.", "Search by keyword, filter by Type, Focus area, or Country / Region, then select Read case study."],
            ["Data Centre", "Explore project data, progress metrics, dashboards, maps, and project detail pages.", "Review charts and tables, select a project title, then use Back to Data Centre."],
            ["Directory", "Find experts by name, organization, role, profile text, or country.", "Search, filter by country, select View Profile, and use pagination when needed."],
            ["Forum", "Enter community discussion spaces and forum topic areas.", "Open a forum space, read topics, and participate if your account has access."],
        ],
        [1.25, 2.55, 2.7],
    )

    add_heading(doc, "2. Getting Started", 1)
    add_heading(doc, "Register for an account", 2)
    add_paragraph(
        doc,
        "Create an account when you need access to member features such as course progress, feedback, certificates, forum participation, or profile-based services.",
    )
    add_number(doc, "Select Sign in in the top-right area of the header.")
    add_number(doc, "On the Login page, select Register.")
    add_number(doc, "Complete the registration form, including username, name, email address, password, and requested profile fields such as sex, age, organization type, and country.")
    add_number(doc, "Submit the form and follow any confirmation instructions shown by the website.")
    add_screenshot(doc, "11-login-register-link.png", "Login page with Register and Forgot your password links.")
    add_screenshot(doc, "12-register-page.png", "Registration page where new users create an account and complete profile fields.")

    add_heading(doc, "Sign in or log out", 2)
    add_number(doc, "Select Sign in in the top-right area of the header.")
    add_number(doc, "Enter your account credentials.")
    add_number(doc, "After signing in, the header changes to Log out. Select Log out when you are finished on a shared computer.")
    add_heading(doc, "Change language", 2)
    add_paragraph(
        doc,
        "Use the language selector in the header. It shows the current language code, such as EN, next to a flag. Select the dropdown and choose the language version you want to view.",
    )
    add_note(
        doc,
        "Translation note",
        "Some content depends on whether a translated version has been created for that page or item. If a specific item is missing in another language, switch back to English or contact the site team.",
    )
    add_screenshot(doc, "10-mobile-menu.png", "Mobile menu with the same primary website sections available from the compact header.")

    add_heading(doc, "3. Homepage", 1)
    add_paragraph(
        doc,
        "The homepage is designed as the main entry point. It combines introduction content with direct paths into the major website sections.",
    )
    add_screenshot(doc, "01-homepage-header.png", "Homepage header and hero area with the main navigation, language selector, and sign-in button.")
    add_bullet(doc, "Hero actions: use the primary and secondary buttons near the top of the page for the most important site pathways.")
    add_bullet(doc, "Search panel: search resources, courses, case studies, experts, webinars, and Data Centre records.")
    add_bullet(doc, "Quick links: open common areas such as learning, resources, data, and directory content.")
    add_bullet(doc, "Featured cards: open highlighted case studies, resources, courses, and events.")
    add_bullet(doc, "Data card: review project progress metrics and open the Data Centre.")

    add_heading(doc, "4. Search", 1)
    add_paragraph(
        doc,
        "The homepage search is the broadest search tool. It returns matching website content and is intended for cross-site discovery.",
    )
    add_number(doc, "Go to the homepage search field.")
    add_number(doc, "Enter a clear keyword or phrase, for example a project name, resource topic, expert name, or course title.")
    add_number(doc, "Select Search.")
    add_number(doc, "Open the result title that best matches your need.")
    add_screenshot(doc, "02-homepage-search-quick-links.png", "Homepage search panel for searching across resources, courses, case studies, experts, webinars, and Data Centre records.")
    add_table(
        doc,
        ["Search target", "Recommended search terms"],
        [
            ["Data Centre projects", "Project title, project type, project status, country, organization, funding source, implementing party, or milestone output."],
            ["Resources", "Resource title, type, topic, source organization, country / region, or publication year."],
            ["Case studies", "Case study title, type, focus area, country / region, lead organization, or related project."],
            ["Experts", "Name, organization, job title, country, role, languages, or profile text."],
            ["Courses", "Course title, course topic, or learning area."],
        ],
        [1.8, 4.7],
    )
    add_note(
        doc,
        "Search tip",
        "If an exact phrase does not return the item you expect, try a shorter phrase or one distinctive word from the title. Then use the result type and excerpt to confirm the correct item.",
    )

    add_heading(doc, "5. Learn", 1)
    add_paragraph(
        doc,
        "The Learn area hosts self-paced courses. Some courses include SCORM packages that track completion and unlock feedback and certificates.",
    )
    add_screenshot(doc, "03-learn-page.png", "Learn landing page where users can access self-paced courses.")
    add_heading(doc, "Open and complete a course", 2)
    add_number(doc, "Select Learn from the main navigation, or open a featured course from the homepage.")
    add_number(doc, "Select the course title or Open course.")
    add_number(doc, "Work through the course content in order. Course progress appears in the course player sidebar.")
    add_number(doc, "When the course reaches verified completion, use the Course feedback button.")
    add_screenshot(doc, "04-course-page.png", "Course detail page where learners open course content and later access feedback and certificates.")
    add_heading(doc, "Submit feedback and view certificate", 2)
    add_number(doc, "Select Submit feedback survey or the feedback button shown below the course.")
    add_number(doc, "If the course is complete, the feedback form opens.")
    add_number(doc, "Complete and submit the survey.")
    add_number(doc, "After feedback is submitted, the certificate becomes available on the course page.")
    add_note(
        doc,
        "Completion gate",
        "The feedback survey is only accepted after the course has a verified completed or passed status. A score display alone may not be enough if the course package has not reported completion.",
    )

    add_heading(doc, "6. Resources", 1)
    add_paragraph(
        doc,
        "Use Resources to find publications, toolkits, templates, learning resources, multimedia resources, and related materials.",
    )
    add_screenshot(doc, "05-resources-filters.png", "Resources library with search, filter controls, resource cards, and View details buttons.")
    add_number(doc, "Select Resources in the main navigation.")
    add_number(doc, "Enter a keyword in Search resources if you know what you are looking for.")
    add_number(doc, "Use Type, Topic, and Country / Region filters to narrow the list.")
    add_number(doc, "Select View details on a resource card to open the resource.")
    add_number(doc, "Select Reset to clear the filters and return to the full library.")
    add_table(
        doc,
        ["Filter", "When to use it"],
        [
            ["Search", "Use for a title, organization, subject, or phrase."],
            ["Type", "Use for categories such as publications, toolkits, templates, learning resources, or multimedia."],
            ["Topic", "Use when browsing by subject matter."],
            ["Country / Region", "Use when looking for geographically relevant material."],
        ],
        [1.65, 4.85],
    )

    add_heading(doc, "7. Case Studies", 1)
    add_paragraph(
        doc,
        "Case Studies provide examples and stories from implementation work. The library uses the same search-and-filter pattern as Resources.",
    )
    add_screenshot(doc, "06-case-studies-filters.png", "Case Studies library with search, filter controls, and Read case study buttons.")
    add_number(doc, "Select Case Studies from the main navigation.")
    add_number(doc, "Search by keyword or use the Type, Focus area, and Country / Region filters.")
    add_number(doc, "Select Read case study to open the full item.")
    add_number(doc, "Use Reset to clear filters.")
    add_note(
        doc,
        "Related content",
        "Some case studies may reference Data Centre projects. When a related project link is available, open it to view project details, progress, outputs, and indicators.",
    )

    add_heading(doc, "8. Data Centre", 1)
    add_paragraph(
        doc,
        "The Data Centre presents project and programme information in dashboard, card, chart, map, and table formats.",
    )
    add_screenshot(doc, "07-data-centre.png", "Data Centre dashboard entry point for stories, project data, and programme metrics.")
    add_bullet(doc, "Dashboard metrics summarize project progress and participation data.")
    add_bullet(doc, "Maps and charts show distribution by territory or project grouping.")
    add_bullet(doc, "Project cards and tables provide direct access to individual project records.")
    add_bullet(doc, "Project detail pages include fields such as country, project type, status, percent complete, implementing party, organizations, funding source, gender responsiveness, people affected, and resilience-related scores.")
    add_heading(doc, "Open a project detail page", 2)
    add_number(doc, "Select Data Centre in the main navigation.")
    add_number(doc, "Find the project in a card, table, or search result.")
    add_number(doc, "Select the project title.")
    add_number(doc, "Review the project summary, metadata, and milestone output.")
    add_number(doc, "Select Back to Data Centre to return to the dashboard.")

    add_heading(doc, "9. Directory", 1)
    add_paragraph(
        doc,
        "The Directory lists experts and partner profiles. It supports keyword search and country filtering.",
    )
    add_screenshot(doc, "08-directory-filters.png", "Expert Directory listing with searchable expert profile cards and country filtering.")
    add_number(doc, "Select Directory from the main navigation.")
    add_number(doc, "Use the search field for a name, organization, role, or profile text.")
    add_number(doc, "Use the country filter if you need experts connected to a specific country.")
    add_number(doc, "Select View Profile to open the full profile.")
    add_number(doc, "Use pagination to move through additional results.")
    add_note(
        doc,
        "No results",
        "If no experts match your filters, clear the country filter or use a broader search term.",
    )

    add_heading(doc, "10. Forum and Events", 1)
    add_heading(doc, "Forum", 2)
    add_paragraph(
        doc,
        "The Forum area is for community discussion spaces. Access may depend on account permissions. Open a forum space to view available topics and participate where posting is enabled.",
    )
    add_screenshot(doc, "09-forum.png", "Forum landing page showing community discussion entry points and forum statistics.")
    add_heading(doc, "Events and webinars", 2)
    add_paragraph(
        doc,
        "Homepage event cards show upcoming or latest events when available. Select View event to open the event page and review timing and participation details.",
    )

    add_heading(doc, "11. Troubleshooting", 1)
    add_table(
        doc,
        ["Issue", "What to try"],
        [
            ["I cannot find an item through search.", "Try a shorter keyword, use a distinctive title word, or navigate directly to Resources, Case Studies, Data Centre, or Directory and filter there."],
            ["Filters show no results.", "Select Reset, then apply one filter at a time. Broaden the keyword if needed."],
            ["I cannot submit course feedback.", "Return to the course player and make sure the course has reported completion or passed status. Then select the feedback button again."],
            ["I submitted feedback but do not see the certificate.", "Refresh the course page. If it still does not appear, confirm you are signed in with the same account used to complete the course."],
            ["A page appears in the wrong language or is missing translated content.", "Use the language dropdown to switch languages. If the item is not translated yet, use the English version or contact the site team."],
            ["The mobile menu is hidden.", "Use the menu button in the header. Select the GSF Hub logo to return to the homepage."],
        ],
        [2.15, 4.35],
    )

    add_heading(doc, "12. Accessibility and Usability Tips", 1)
    add_bullet(doc, "Use the Skip to content link when navigating by keyboard or screen reader.")
    add_bullet(doc, "Use descriptive search terms instead of very broad terms when looking for a specific item.")
    add_bullet(doc, "Use Reset links after filtering so you know you are looking at the full list again.")
    add_bullet(doc, "Sign out after using the website on a shared or public computer.")
    add_bullet(doc, "Keep the same signed-in account throughout a course so progress, feedback, and certificate status stay connected.")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_manual()
