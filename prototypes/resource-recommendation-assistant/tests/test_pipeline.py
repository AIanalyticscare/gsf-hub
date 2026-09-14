from __future__ import annotations

import csv
import json
import sys
import tempfile
import threading
import unittest
import urllib.error
import urllib.request
from datetime import date
from pathlib import Path


PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT / "src"))

from resource_recommender.ingest import ingest_path  # noqa: E402
from resource_recommender.recommender import ResourceRecommender  # noqa: E402
from resource_recommender.server import RecommendationHandler, RecommendationServer  # noqa: E402


class RecommendationPipelineTests(unittest.TestCase):
    def test_eleuthera_tourism_query_ranks_resilience_grant_first(self) -> None:
        library = ingest_path(PROJECT_ROOT / "data")
        recommender = ResourceRecommender(library, today=date(2026, 8, 3))
        need = recommender.understand(
            "I manage a small tourism business in Eleuthera and need support for hurricane resilience."
        )
        result = recommender.recommend(need)

        self.assertEqual("Eleuthera", result.query.location)
        self.assertEqual("Tourism", result.query.sector)
        self.assertEqual("Small Business Resilience Grant", result.recommendations[0].resource)
        self.assertGreater(result.recommendations[0].fit_score, result.recommendations[1].fit_score)
        self.assertEqual("demo_resources.csv", result.recommendations[0].sources[0]["document"])
        self.assertEqual(2, result.recommendations[0].sources[0]["row"])
        self.assertTrue(any("demo_unverified" in item for item in result.recommendations[0].uncertainties))

    def test_expired_deadline_is_not_presented_as_current(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "resources.csv"
            with path.open("w", encoding="utf-8", newline="") as handle:
                writer = csv.DictWriter(
                    handle,
                    fieldnames=["resource_name", "description", "eligible_islands", "application_deadline"],
                )
                writer.writeheader()
                writer.writerow(
                    {
                        "resource_name": "Expired Grant",
                        "description": "Hurricane resilience grant",
                        "eligible_islands": "Eleuthera",
                        "application_deadline": "2025-01-01",
                    }
                )
            recommender = ResourceRecommender(ingest_path(path), today=date(2026, 8, 3))
            result = recommender.recommend(recommender.understand("Eleuthera hurricane grant"))
            item = result.recommendations[0]
            self.assertEqual("unlikely_or_ineligible", item.eligibility_status)
            self.assertTrue(any("has passed" in barrier for barrier in item.possible_barriers))

    def test_reads_pdf_word_excel_and_csv_with_locators(self) -> None:
        try:
            from docx import Document
            from openpyxl import Workbook
            from reportlab.pdfgen import canvas
        except ImportError as exc:  # pragma: no cover - local environment guard
            self.skipTest(str(exc))

        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            (root / "one.csv").write_text(
                "resource_name,description\nCSV Resource,Storm planning support\n", encoding="utf-8"
            )

            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Programmes"
            sheet.append(["resource_name", "description"])
            sheet.append(["Excel Resource", "Business training"])
            workbook.save(root / "two.xlsx")

            document = Document()
            document.add_paragraph("Word guidance for coastal restoration.")
            document.save(root / "three.docx")

            pdf = canvas.Canvas(str(root / "four.pdf"))
            pdf.drawString(72, 720, "PDF handbook for biodiversity grants.")
            pdf.save()

            library = ingest_path(root)
            names = {resource.resource_name for resource in library.resources}
            self.assertTrue({"CSV Resource", "Excel Resource", "three", "four"}.issubset(names))
            pdf_resource = next(resource for resource in library.resources if resource.resource_name == "four")
            self.assertEqual(1, pdf_resource.evidence[0].source.page)
            word_resource = next(resource for resource in library.resources if resource.resource_name == "three")
            self.assertEqual(1, word_resource.evidence[0].source.paragraph)

    def test_reads_gsf_wordpress_transfer_resource_posts(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "transfer.json"
            path.write_text(
                json.dumps(
                    {
                        "format": "gsf-hub-transfer",
                        "posts": [
                            {
                                "post_type": "gsf_resource",
                                "post_status": "publish",
                                "post_title": "WordPress Toolkit",
                                "post_excerpt": "A toolkit for hurricane planning.",
                                "post_modified": "2026-07-01 12:00:00",
                                "url": "http://example.test/gsf-resource/wordpress-toolkit/",
                                "meta": {
                                    "resource_type": ["Toolkit"],
                                    "country_region": ["Bahamas"],
                                    "topic": ["Climate resilience"],
                                },
                            },
                            {
                                "post_type": "gsf_course",
                                "post_status": "publish",
                                "post_title": "Resilience Course",
                                "post_excerpt": "A course about climate resilience planning.",
                                "post_modified": "2026-07-02 12:00:00",
                                "url": "http://example.test/learn/resilience-course/",
                                "meta": {},
                                "topics": ["Climate resilience"],
                            },
                            {
                                "post_type": "topic",
                                "post_status": "publish",
                                "post_title": "Share adaptation lessons",
                                "post_content": "A public discussion about practical adaptation lessons.",
                                "post_modified": "2026-07-03 12:00:00",
                                "url": "http://example.test/forums/topic/adaptation-lessons/",
                                "meta": {},
                            },
                            {"post_type": "page", "post_title": "Ignore me"},
                        ],
                    }
                ),
                encoding="utf-8",
            )
            library = ingest_path(path)
            self.assertEqual(3, len(library.resources))
            by_name = {resource.resource_name: resource for resource in library.resources}
            self.assertEqual(["Bahamas"], by_name["WordPress Toolkit"].eligible_islands)
            self.assertEqual(["Climate resilience"], by_name["WordPress Toolkit"].needs_addressed)
            self.assertEqual("Learn course", by_name["Resilience Course"].resource_type)
            self.assertEqual("Forum discussion", by_name["Share adaptation lessons"].resource_type)
            self.assertEqual(
                "http://example.test/learn/resilience-course/",
                by_name["Resilience Course"].evidence[0].source.url,
            )

    def test_nctf_query_maps_multiple_hub_paths(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "hub.csv"
            path.write_text(
                "resource_name,description,resource_type,needs_addressed,updated_at\n"
                'NCTF Staff Indicator,"Tracks NCTF and CBF staff training.",Data Centre indicator,"NCTF training|Gender concepts",2026-07-01\n'
                'Conservation Fund Course,"Learning for CTF and NCTF staff.",Learn course,"Conservation fund management|Capacity building",2026-07-01\n'
                'NCTF Governance Toolkit,"Practical tools for conservation trust funds.",Resource · Toolkit,"Governance|Organizational strengthening",2026-07-01\n'
                'Regional NCTF Project,"A project involving NCTF partners.",Data Centre project,"NCTF partnerships",2026-07-01\n',
                encoding="utf-8",
            )
            recommender = ResourceRecommender(ingest_path(path), today=date(2026, 8, 3))
            result = recommender.recommend(
                recommender.understand(
                    "I work for an NCTF and want to strengthen governance and operational capacity. "
                    "What courses, tools, and Data Centre evidence should I use?"
                ),
                top_k=3,
            )
            types = [item.resource_type for item in result.recommendations]
            self.assertTrue(any(item.startswith("Data Centre") for item in types))
            self.assertIn("Learn course", types)
            self.assertIn("Resource · Toolkit", types)
            self.assertFalse(types[0].startswith("Data Centre"))

    def test_http_api_requires_token_and_returns_citations(self) -> None:
        library = ingest_path(PROJECT_ROOT / "data")
        try:
            server = RecommendationServer(("127.0.0.1", 0), RecommendationHandler)
        except PermissionError as exc:  # Some restricted runners disallow loopback sockets.
            self.skipTest(str(exc))
        server.recommender = ResourceRecommender(library, today=date(2026, 8, 3))
        server.api_token = "test-token"
        thread = threading.Thread(target=server.serve_forever, daemon=True)
        thread.start()
        endpoint = f"http://127.0.0.1:{server.server_port}/recommend"
        body = json.dumps(
            {
                "situation": (
                    "I manage a small tourism business in Eleuthera and need support for hurricane resilience."
                )
            }
        ).encode()
        try:
            with self.assertRaises(urllib.error.HTTPError) as context:
                urllib.request.urlopen(
                    urllib.request.Request(endpoint, data=body, headers={"Content-Type": "application/json"}),
                    timeout=3,
                )
            self.assertEqual(401, context.exception.code)

            request = urllib.request.Request(
                endpoint,
                data=body,
                headers={"Content-Type": "application/json", "Authorization": "Bearer test-token"},
            )
            with urllib.request.urlopen(request, timeout=3) as response:
                payload = json.load(response)
            self.assertEqual("Small Business Resilience Grant", payload["recommendations"][0]["resource"])
            self.assertTrue(payload["recommendations"][0]["sources"])
        finally:
            server.shutdown()
            server.server_close()
            thread.join(timeout=3)


if __name__ == "__main__":
    unittest.main()
