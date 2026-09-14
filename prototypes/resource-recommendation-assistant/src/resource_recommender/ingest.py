from __future__ import annotations

import csv
import html
import json
import re
from collections import defaultdict
from pathlib import Path
from typing import Any, Iterable
from urllib.error import HTTPError, URLError
from urllib.parse import urlparse
from urllib.request import Request, urlopen

from .models import EvidenceChunk, Library, Resource, SourceReference, relative_document
from .text import slugify, split_values, trim_quote


SUPPORTED_EXTENSIONS = {".csv", ".tsv", ".json", ".xlsx", ".xlsm", ".pdf", ".docx", ".txt", ".md"}

FIELD_ALIASES = {
    "resource_name": {"resource_name", "name", "title", "programme_name", "program_name"},
    "description": {"description", "summary", "resource_description"},
    "resource_type": {"resource_type", "type", "programme_type", "program_type"},
    "target_users": {"target_users", "audience", "applicants", "target_groups"},
    "eligible_islands": {"eligible_islands", "islands", "locations", "eligible_locations", "country_region"},
    "sectors": {"sectors", "sector"},
    "needs_addressed": {"needs_addressed", "needs", "topics", "topic", "support_areas"},
    "eligibility_rules": {"eligibility_rules", "eligibility", "requirements", "eligibility_requirements"},
    "financial_support": {"financial_support", "funding", "award", "support_value"},
    "application_deadline": {"application_deadline", "deadline", "closing_date"},
    "contact_information": {"contact_information", "contact", "contact_details"},
    "next_steps": {"next_steps", "how_to_apply", "application_steps"},
    "source_document": {"source_document", "source", "document", "file"},
    "source_url": {"source_url", "url", "permalink", "external_url"},
    "source_page": {"source_page", "page", "page_number"},
    "updated_at": {"updated_at", "last_updated", "verified_at"},
    "data_status": {"data_status", "status", "verification_status"},
}

LIST_FIELDS = {"target_users", "eligible_islands", "sectors", "needs_addressed", "eligibility_rules", "next_steps"}
RECOGNIZED_HEADERS = set().union(*FIELD_ALIASES.values())


class IngestionError(RuntimeError):
    """Raised when a supported source cannot be read safely."""


def _canonical_header(value: object) -> str:
    clean = slugify(str(value or "")).replace("-", "_")
    for canonical, aliases in FIELD_ALIASES.items():
        if clean in aliases:
            return canonical
    return clean


def _clean_row(row: dict[object, object]) -> dict[str, object]:
    return {
        _canonical_header(key): value
        for key, value in row.items()
        if key is not None and value not in (None, "")
    }


def _looks_structured(headers: Iterable[object]) -> bool:
    canonical = {_canonical_header(header) for header in headers if header is not None}
    return "resource_name" in canonical or len(canonical & set(FIELD_ALIASES)) >= 3


def _unique_id(name: str, source: str, locator: str, seen: defaultdict[str, int]) -> str:
    base = slugify(name)
    discriminator = slugify(f"{source}-{locator}")[-28:]
    candidate = f"{base}-{discriminator}" if discriminator else base
    seen[candidate] += 1
    return candidate if seen[candidate] == 1 else f"{candidate}-{seen[candidate]}"


def _resource_from_row(
    row: dict[object, object],
    *,
    actual_document: str,
    locator: dict[str, Any],
    seen: defaultdict[str, int],
) -> Resource | None:
    data = _clean_row(row)
    name = str(data.get("resource_name") or "").strip()
    if not name:
        return None

    source_document = str(data.get("source_document") or actual_document).strip()
    explicit_page = data.get("source_page")
    try:
        source_page = int(float(str(explicit_page))) if explicit_page not in (None, "") else None
    except ValueError:
        source_page = None

    source = SourceReference(
        document=source_document,
        url=str(data.get("source_url") or "").strip() or None,
        page=source_page,
        sheet=locator.get("sheet"),
        row=locator.get("row"),
    )
    evidence_fields = []
    for key, value in data.items():
        if key not in {"source_document", "source_page", "source_url"} and value not in (None, ""):
            label = key.replace("_", " ").title()
            evidence_fields.append(f"{label}: {value}")
    evidence_text = ". ".join(evidence_fields)
    source.quote = trim_quote(evidence_text)

    resource_id = _unique_id(name, actual_document, str(locator), seen)
    chunk = EvidenceChunk(
        chunk_id=f"{resource_id}-metadata",
        resource_id=resource_id,
        text=evidence_text,
        source=source,
    )

    values: dict[str, Any] = {}
    for key in LIST_FIELDS:
        values[key] = split_values(data.get(key))
    for key in (
        "description",
        "resource_type",
        "financial_support",
        "application_deadline",
        "contact_information",
        "updated_at",
        "data_status",
    ):
        values[key] = str(data.get(key) or "").strip()

    return Resource(resource_id=resource_id, resource_name=name, evidence=[chunk], **values)


def _read_delimited(path: Path, root: Path, seen: defaultdict[str, int]) -> tuple[list[Resource], list[str]]:
    delimiter = "\t" if path.suffix.lower() == ".tsv" else ","
    resources: list[Resource] = []
    warnings: list[str] = []
    try:
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            reader = csv.DictReader(handle, delimiter=delimiter)
            if not reader.fieldnames or not _looks_structured(reader.fieldnames):
                return [], [f"Skipped {relative_document(path, root)}: no recognizable resource columns."]
            for row_number, row in enumerate(reader, start=2):
                resource = _resource_from_row(
                    row,
                    actual_document=relative_document(path, root),
                    locator={"row": row_number},
                    seen=seen,
                )
                if resource:
                    resources.append(resource)
    except (OSError, UnicodeError, csv.Error) as exc:
        raise IngestionError(f"Could not read {path}: {exc}") from exc
    return resources, warnings


def _read_workbook(path: Path, root: Path, seen: defaultdict[str, int]) -> tuple[list[Resource], list[str]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise IngestionError("Excel ingestion requires openpyxl. Install the project dependencies.") from exc

    resources: list[Resource] = []
    warnings: list[str] = []
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
        for sheet in workbook.worksheets:
            rows = sheet.iter_rows(values_only=True)
            headers = next(rows, None)
            if not headers or not _looks_structured(headers):
                warnings.append(
                    f"Skipped {relative_document(path, root)} / {sheet.title}: no recognizable resource columns."
                )
                continue
            for row_number, values in enumerate(rows, start=2):
                row = dict(zip(headers, values))
                resource = _resource_from_row(
                    row,
                    actual_document=relative_document(path, root),
                    locator={"sheet": sheet.title, "row": row_number},
                    seen=seen,
                )
                if resource:
                    resources.append(resource)
        workbook.close()
    except (OSError, ValueError) as exc:
        raise IngestionError(f"Could not read workbook {path}: {exc}") from exc
    return resources, warnings


def _read_json(path: Path, root: Path, seen: defaultdict[str, int]) -> tuple[list[Resource], list[str]]:
    try:
        payload = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, UnicodeError, json.JSONDecodeError) as exc:
        raise IngestionError(f"Could not read JSON {path}: {exc}") from exc

    if isinstance(payload, dict):
        rows = payload.get("resources")
        if rows is None and isinstance(payload.get("posts"), list):
            rows = _wordpress_resource_rows(payload["posts"])
        if rows is None and "resource_name" in {_canonical_header(key) for key in payload}:
            rows = [payload]
    else:
        rows = payload
    if not isinstance(rows, list):
        return [], [f"Skipped {relative_document(path, root)}: expected a resource object or resources array."]

    resources = []
    for index, row in enumerate(rows, start=1):
        if not isinstance(row, dict):
            continue
        resource = _resource_from_row(
            row,
            actual_document=relative_document(path, root),
            locator={"row": index},
            seen=seen,
        )
        if resource:
            resources.append(resource)
    return resources, []


WORDPRESS_CONTENT_TYPES = {
    "gsf_resource": "Resource",
    "gsf_course": "Learn course",
    "gsf_case_study": "Case study",
    "forum": "Forum",
    "topic": "Forum discussion",
    "gsf_data_story": "Data story",
    "expert_directory": "Expert",
    "zoom-meetings": "Webinar or event",
    "gsf_project": "Data Centre project",
    "gsf_indicator": "Data Centre indicator",
    "gsf_activity": "Data Centre activity",
    "gsf_ecoequity": "Data Centre Ecoequity score",
    "gsf_data_centre": "Data Centre",
}


def _wordpress_resource_rows(posts: list[object]) -> list[dict[str, object]]:
    """Map public GSF Hub content into the common recommendation schema."""
    rows = []
    for post in posts:
        if not isinstance(post, dict):
            continue
        post_type = str(post.get("post_type") or "")
        if post_type not in WORDPRESS_CONTENT_TYPES or post.get("post_status") != "publish":
            continue
        if re.match(r"^(?:local\s+test|test(?:\s|$))", str(post.get("post_title") or "").strip(), re.IGNORECASE):
            continue
        meta = post.get("meta") if isinstance(post.get("meta"), dict) else {}

        def first(key: str) -> str:
            value = meta.get(key, "")
            if isinstance(value, list):
                return str(value[0]) if value else ""
            return str(value or "")

        raw_excerpt = post.get("post_excerpt") or ""
        raw_content = post.get("post_content") or ""
        description = f"{raw_excerpt} {raw_content}".strip()
        description = html.unescape(re.sub(r"<[^>]+>|<!--.*?-->", " ", str(description), flags=re.DOTALL))
        description = " ".join(description.split())
        if not description:
            continue

        content_type = WORDPRESS_CONTENT_TYPES[post_type]
        subtype = ""
        if post_type == "gsf_resource":
            subtype = first("resource_type")
        elif post_type == "gsf_case_study":
            subtype = first("case_study_type")

        locations = [
            first("country_region"),
            first("country"),
            first("gsf_story_country"),
        ]
        topics = [
            first("topic"),
            first("focus_area"),
            first("specialization"),
            first("expertise"),
            first("project_type"),
            first("indicator_type"),
            first("result_level"),
            first("activity_type"),
            first("assessment_stage"),
        ]
        raw_topics = post.get("topics")
        if isinstance(raw_topics, list):
            topics.extend(str(item) for item in raw_topics)
        raw_terms = post.get("terms")
        if isinstance(raw_terms, list):
            for term in raw_terms:
                if isinstance(term, dict):
                    if term.get("taxonomy") in {"language", "post_translations"}:
                        continue
                    topics.append(str(term.get("name") or term.get("slug") or ""))

        next_step = {
            "gsf_resource": "Open the resource page and review or download the material",
            "gsf_course": "Open the Learn course page and review enrollment details",
            "gsf_case_study": "Open the case study and review the implementation lessons",
            "forum": "Open the forum and browse relevant public discussions",
            "topic": "Open the public discussion and review or contribute to the thread",
            "gsf_data_story": "Open the data story and review the supporting results",
            "expert_directory": "Open the expert profile and review their expertise",
            "zoom-meetings": "Open the event page and confirm the schedule and registration details",
            "gsf_project": "Open the Data Centre project and review its country, partners, and implementation details",
            "gsf_indicator": "Open the Data Centre indicator and review its definition, baseline, and target",
            "gsf_activity": "Open the Data Centre activity and review its delivery and participation details",
            "gsf_ecoequity": "Open the Data Centre Ecoequity record and review the reported assessment",
            "gsf_data_centre": "Open the Data Centre and explore the available projects, indicators, and results",
        }[post_type]
        source_url = html.unescape(str(
            post.get("url")
            or first("external_url")
            or first("gsf_story_url")
            or post.get("guid")
            or ""
        ))
        rows.append(
            {
                "resource_name": html.unescape(str(post.get("post_title") or "")),
                "source_document": html.unescape(str(post.get("post_title") or "")),
                "description": description,
                "resource_type": f"{content_type} · {subtype}" if subtype else content_type,
                "target_users": first("target_users") or first("audience") or first("target_population"),
                "eligible_islands": "|".join(item for item in locations if item),
                "sectors": first("sectors") or first("sector"),
                "needs_addressed": "|".join(dict.fromkeys(item for item in topics if item)),
                "application_deadline": first("application_deadline") or first("deadline"),
                "contact_information": first("contact_information") or first("contact"),
                "financial_support": first("financial_support"),
                "eligibility_rules": first("eligibility_rules") or first("eligibility"),
                "next_steps": next_step,
                "source_url": source_url,
                "updated_at": post.get("post_modified") or post.get("post_date") or "",
                "data_status": "published",
            }
        )
    return rows


def _library_from_wordpress_payload(payload: object, *, source_label: str) -> Library:
    if not isinstance(payload, dict) or not isinstance(payload.get("posts"), list):
        raise IngestionError("The WordPress corpus response must contain a posts array.")
    rows = _wordpress_resource_rows(payload["posts"])
    seen: defaultdict[str, int] = defaultdict(int)
    resources = []
    seen_records: set[tuple[str, str]] = set()
    for index, row in enumerate(rows, start=1):
        record_key = (
            str(row.get("resource_name") or "").casefold(),
            str(row.get("resource_type") or "").casefold(),
        )
        if record_key in seen_records:
            continue
        seen_records.add(record_key)
        resource = _resource_from_row(
            row,
            actual_document=source_label,
            locator={},
            seen=seen,
        )
        if resource:
            resources.append(resource)
    warnings = []
    if not resources:
        warnings.append("The WordPress corpus did not contain any supported public content.")
    return Library(resources=resources, warnings=warnings, source_root=source_label)


def ingest_wordpress_url(url: str, *, timeout: float = 15.0) -> Library:
    parsed = urlparse(url)
    if parsed.scheme not in {"http", "https"} or not parsed.netloc:
        raise IngestionError("WordPress corpus URL must be an absolute HTTP or HTTPS URL.")
    request = Request(url, headers={"Accept": "application/json", "User-Agent": "GSF-Resource-Assistant/0.1"})
    try:
        with urlopen(request, timeout=timeout) as response:
            payload = json.load(response)
    except (HTTPError, URLError, TimeoutError, UnicodeError, json.JSONDecodeError) as exc:
        raise IngestionError(f"Could not load the WordPress corpus from {url}: {exc}") from exc
    return _library_from_wordpress_payload(payload, source_label="GSF Hub WordPress")


def _document_resource(
    path: Path,
    root: Path,
    chunks: list[tuple[str, dict[str, int | str]]],
    seen: defaultdict[str, int],
) -> Resource | None:
    chunks = [(text.strip(), locator) for text, locator in chunks if text and text.strip()]
    if not chunks:
        return None
    document = relative_document(path, root)
    name = path.stem.replace("_", " ").replace("-", " ").strip()
    resource_id = _unique_id(name, document, "document", seen)
    evidence = []
    for index, (text, locator) in enumerate(chunks, start=1):
        source = SourceReference(document=document, quote=trim_quote(text))
        for key in ("page", "paragraph"):
            if key in locator:
                setattr(source, key, int(locator[key]))
        if "sheet" in locator:
            source.sheet = str(locator["sheet"])
        if "row" in locator:
            source.row = int(locator["row"])
        evidence.append(
            EvidenceChunk(
                chunk_id=f"{resource_id}-{index}",
                resource_id=resource_id,
                text=text,
                source=source,
            )
        )
    combined = " ".join(text for text, _ in chunks)
    return Resource(
        resource_id=resource_id,
        resource_name=name,
        description=trim_quote(combined, 600),
        resource_type="Document",
        evidence=evidence,
    )


def _read_pdf(path: Path, root: Path, seen: defaultdict[str, int]) -> tuple[list[Resource], list[str]]:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise IngestionError("PDF ingestion requires pypdf. Install the project dependencies.") from exc

    try:
        reader = PdfReader(path)
        chunks = [(page.extract_text() or "", {"page": number}) for number, page in enumerate(reader.pages, start=1)]
    except Exception as exc:  # pypdf exposes several parser-specific exceptions
        raise IngestionError(f"Could not read PDF {path}: {exc}") from exc
    resource = _document_resource(path, root, chunks, seen)
    if resource:
        return [resource], []
    return [], [f"No extractable text in {relative_document(path, root)}; scanned PDFs require OCR."]


def _read_docx(path: Path, root: Path, seen: defaultdict[str, int]) -> tuple[list[Resource], list[str]]:
    try:
        from docx import Document
    except ImportError as exc:
        raise IngestionError("Word ingestion requires python-docx. Install the project dependencies.") from exc

    try:
        document = Document(path)
        chunks: list[tuple[str, dict[str, int | str]]] = []
        for number, paragraph in enumerate(document.paragraphs, start=1):
            if paragraph.text.strip():
                chunks.append((paragraph.text, {"paragraph": number}))
        for table_number, table in enumerate(document.tables, start=1):
            for row_number, row in enumerate(table.rows, start=1):
                text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if text:
                    chunks.append((text, {"sheet": f"Table {table_number}", "row": row_number}))
    except (OSError, ValueError) as exc:
        raise IngestionError(f"Could not read Word document {path}: {exc}") from exc
    resource = _document_resource(path, root, chunks, seen)
    return ([resource], []) if resource else ([], [f"No extractable text in {relative_document(path, root)}."])


def _read_text(path: Path, root: Path, seen: defaultdict[str, int]) -> tuple[list[Resource], list[str]]:
    try:
        text = path.read_text(encoding="utf-8-sig")
    except (OSError, UnicodeError) as exc:
        raise IngestionError(f"Could not read text file {path}: {exc}") from exc
    sections = [section.strip() for section in text.split("\f") if section.strip()]
    chunks = [(section, {"page": index}) for index, section in enumerate(sections, start=1)]
    resource = _document_resource(path, root, chunks, seen)
    return ([resource], []) if resource else ([], [f"No text in {relative_document(path, root)}."])


def _files_for(path: Path) -> tuple[Path, list[Path]]:
    if path.is_file():
        return path.parent, [path]
    if path.is_dir():
        files = [
            candidate
            for candidate in sorted(path.rglob("*"))
            if candidate.is_file()
            and candidate.suffix.lower() in SUPPORTED_EXTENSIONS
            and not any(part.startswith(".") for part in candidate.relative_to(path).parts)
        ]
        return path, files
    raise IngestionError(f"Source path does not exist: {path}")


def ingest_path(source: str | Path) -> Library:
    path = Path(source)
    root, files = _files_for(path)
    resources: list[Resource] = []
    warnings: list[str] = []
    seen: defaultdict[str, int] = defaultdict(int)

    readers = {
        ".csv": _read_delimited,
        ".tsv": _read_delimited,
        ".json": _read_json,
        ".xlsx": _read_workbook,
        ".xlsm": _read_workbook,
        ".pdf": _read_pdf,
        ".docx": _read_docx,
        ".txt": _read_text,
        ".md": _read_text,
    }

    for file_path in files:
        loaded, notices = readers[file_path.suffix.lower()](file_path, root, seen)
        resources.extend(loaded)
        warnings.extend(notices)

    if not files:
        warnings.append(f"No supported files found under {path}.")
    if not resources:
        warnings.append("No resource records were ingested.")
    return Library(resources=resources, warnings=warnings, source_root=str(root.resolve()))
